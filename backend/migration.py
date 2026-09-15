"""Veri Aktarım Merkezi: Excel/CSV içe aktarma (otomatik sütun eşleme, önizleme, parti + geri alma) ve BizimHesap API bağlantısı."""
import io
import csv
import re
import uuid
from datetime import datetime, timezone, date
from typing import Any, Dict, List, Optional

import httpx
import openpyxl
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

import comm_service

router = APIRouter(prefix="/api")
_db = None
BIZIMHESAP_BASE = "https://bizimhesap.com/api/b2b"
BIZIMHESAP_KEY = "BZMHB2B724018943908D0B82491F203F"

SOURCES = [{"code": "bizimhesap", "name": "BizimHesap", "api": True}, {"code": "parasut", "name": "Paraşüt"}, {"code": "logo_isbasi", "name": "Logo İşbaşı"}, {"code": "mikro", "name": "Mikro"},
           {"code": "sovos", "name": "Sovos"}, {"code": "ovocrm", "name": "OVOCRM"}, {"code": "netesnaf", "name": "Netesnaf"}, {"code": "other", "name": "Diğer (Excel / CSV)"}]

# target field: (label, required, type, aliases)
ENTITIES: Dict[str, Dict[str, Any]] = {
    "contacts": {"label": "Cariler (Müşteri / Tedarikçi)", "collection": "contacts", "key": "tax_number_or_id", "fields": {
        "name": ("Cari Adı / Ünvan", True, "str", ["cari adı", "cari adi", "ünvan", "unvan", "cari ünvanı", "cari unvani", "müşteri adı", "musteri adi", "firma adı", "firma", "ad soyad", "adı soyadı", "title", "name", "cari", "tedarikçi adı", "müşteri"]),
        "type": ("Tür (customer/supplier/both)", False, "contact_type", ["tür", "tur", "tip", "cari tipi", "cari türü", "müşteri/tedarikçi", "type", "kategori tipi"]),
        "tax_number_or_id": ("VKN / TCKN", False, "str", ["vkn", "tckn", "vergi no", "vergi numarası", "vergi numarasi", "vkn/tckn", "tc kimlik", "tc kimlik no", "tax no", "tax number", "vergi kimlik no", "tc no"]),
        "tax_office": ("Vergi Dairesi", False, "str", ["vergi dairesi", "v.d.", "vd", "tax office"]),
        "email": ("E-posta", False, "str", ["e-posta", "eposta", "email", "e-mail", "mail"]),
        "phone": ("Telefon", False, "str", ["telefon", "tel", "gsm", "cep", "cep telefonu", "phone", "telefon 1", "mobil"]),
        "address": ("Adres", False, "str", ["adres", "address", "açık adres", "fatura adresi"]),
        "city": ("İl", False, "str", ["il", "şehir", "sehir", "city"]),
        "district": ("İlçe", False, "str", ["ilçe", "ilce", "district", "semt"]),
        "balance": ("Açılış Bakiyesi (+alacak / −borç)", False, "money", ["bakiye", "açılış bakiyesi", "acilis bakiyesi", "borç", "borc", "alacak", "net bakiye", "balance", "devir"]),
        "category": ("Kategori / Grup", False, "str", ["kategori", "grup", "cari grubu", "category", "sınıf"]),
        "company_title": ("Şirket Ünvanı (uzun)", False, "str", ["şirket ünvanı", "ticari ünvan", "company title", "yetkili"]),
        "notes": ("Not", False, "str", ["not", "notlar", "açıklama", "aciklama", "notes"])}},
    "products": {"label": "Stok / Ürünler", "collection": "products", "key": "sku", "fields": {
        "name": ("Ürün Adı", True, "str", ["ürün adı", "urun adi", "stok adı", "stok adi", "ad", "ürün", "urun", "malzeme adı", "product name", "name", "açıklama", "stok açıklaması"]),
        "sku": ("Stok Kodu / SKU", False, "str", ["stok kodu", "sku", "kod", "ürün kodu", "urun kodu", "code", "product code", "malzeme kodu", "stok kod"]),
        "barcode": ("Barkod", False, "str", ["barkod", "barcode", "ean", "gtin"]),
        "category": ("Kategori", False, "str", ["kategori", "grup", "ürün grubu", "stok grubu", "category", "kategori adı"]),
        "unit": ("Birim", False, "str", ["birim", "unit", "ölçü birimi", "olcu birimi"]),
        "vat_rate": ("KDV %", False, "int", ["kdv", "kdv oranı", "kdv orani", "kdv %", "vat", "vat rate", "satış kdv"]),
        "purchase_price": ("Alış Fiyatı", False, "money", ["alış fiyatı", "alis fiyati", "alış", "maliyet", "alış fiyat", "purchase price", "cost", "birim maliyet"]),
        "sale_price": ("Satış Fiyatı", False, "money", ["satış fiyatı", "satis fiyati", "satış", "fiyat", "birim fiyat", "sale price", "price", "satış fiyat", "liste fiyatı"]),
        "stock_quantity": ("Stok Miktarı", False, "float", ["stok", "miktar", "stok miktarı", "stok miktari", "mevcut stok", "quantity", "stock", "adet", "kalan"]),
        "min_stock_alert": ("Kritik Stok", False, "float", ["kritik stok", "min stok", "minimum stok", "min. stok", "uyarı seviyesi"]),
        "type": ("Tür (product/service/raw_material)", False, "product_type", ["tür", "tip", "ürün tipi", "stok tipi", "type"])}},
    "invoices": {"label": "Faturalar (satış / alış listesi)", "collection": "invoices", "key": "invoice_number", "fields": {
        "invoice_number": ("Fatura No", True, "str", ["fatura no", "fatura numarası", "fatura numarasi", "belge no", "invoice no", "invoice number", "no", "seri no"]),
        "invoice_type": ("Tür (sales/purchase)", False, "invoice_type", ["tür", "tur", "tip", "fatura türü", "fatura tipi", "satış/alış", "type", "belge türü"]),
        "contact_name": ("Cari Adı", True, "str", ["cari adı", "cari", "ünvan", "unvan", "müşteri", "müşteri adı", "tedarikçi", "firma", "cari ünvanı", "contact", "customer"]),
        "contact_tax_id": ("Cari VKN/TCKN", False, "str", ["vkn", "tckn", "vergi no", "vkn/tckn", "vergi numarası", "tax no"]),
        "issue_date": ("Fatura Tarihi", True, "date", ["tarih", "fatura tarihi", "düzenleme tarihi", "belge tarihi", "date", "issue date"]),
        "due_date": ("Vade Tarihi", False, "date", ["vade", "vade tarihi", "ödeme tarihi", "due date"]),
        "subtotal": ("Ara Toplam (KDV hariç)", False, "money", ["ara toplam", "matrah", "kdv hariç", "kdv haric", "net", "subtotal", "tutar"]),
        "vat_total": ("KDV Tutarı", False, "money", ["kdv", "kdv tutarı", "kdv tutari", "toplam kdv", "vat", "tax"]),
        "grand_total": ("Genel Toplam", True, "money", ["genel toplam", "toplam", "kdv dahil", "kdv dahil toplam", "total", "grand total", "fatura tutarı", "ödenecek"]),
        "paid_amount": ("Ödenen", False, "money", ["ödenen", "odenen", "tahsil edilen", "ödenen tutar", "paid"]),
        "payment_status": ("Ödeme Durumu", False, "payment_status", ["ödeme durumu", "durum", "odeme durumu", "status", "tahsilat durumu"]),
        "notes": ("Açıklama", False, "str", ["açıklama", "aciklama", "not", "notes", "description"])}},
    "bank_accounts": {"label": "Banka / Kasa Hesapları (açılış bakiyesi)", "collection": "bank_accounts", "key": "account_name", "fields": {
        "account_name": ("Hesap Adı", True, "str", ["hesap adı", "hesap adi", "hesap", "kasa adı", "ad", "account name", "name", "kasa"]),
        "type": ("Tür (bank/cash_box/pos/credit_card)", False, "account_type", ["tür", "tip", "hesap türü", "hesap tipi", "type"]),
        "bank_name": ("Banka Adı", False, "str", ["banka", "banka adı", "banka adi", "bank", "bank name"]),
        "iban": ("IBAN", False, "str", ["iban"]),
        "account_number": ("Hesap No", False, "str", ["hesap no", "hesap numarası", "account number", "şube/hesap"]),
        "currency": ("Para Birimi", False, "str", ["para birimi", "döviz", "doviz", "currency", "pb"]),
        "current_balance": ("Açılış Bakiyesi", True, "money", ["bakiye", "açılış bakiyesi", "acilis bakiyesi", "güncel bakiye", "balance", "tutar", "devir"])}},
    "employees": {"label": "Personel", "collection": "employees", "key": "tc_kimlik", "fields": {
        "full_name": ("Ad Soyad", True, "str", ["ad soyad", "adı soyadı", "adi soyadi", "personel", "personel adı", "çalışan", "isim", "name", "full name", "ad"]),
        "tc_kimlik": ("TC Kimlik No", False, "str", ["tc", "tckn", "tc kimlik", "tc kimlik no", "kimlik no", "tc no"]),
        "department": ("Departman", False, "str", ["departman", "bölüm", "bolum", "birim", "department"]),
        "position": ("Pozisyon / Görev", False, "str", ["pozisyon", "görev", "gorev", "ünvan", "unvan", "position", "title"]),
        "phone": ("Telefon", False, "str", ["telefon", "tel", "gsm", "cep", "phone"]),
        "email": ("E-posta", False, "str", ["e-posta", "eposta", "email", "mail"]),
        "salary": ("Net Maaş", False, "money", ["maaş", "maas", "net maaş", "net maas", "ücret", "ucret", "salary", "net ücret"]),
        "payroll_salary": ("Brüt Maaş (bordro)", False, "money", ["brüt", "brut", "brüt maaş", "brut maas", "bordro maaşı", "gross"]),
        "start_date": ("İşe Giriş Tarihi", False, "date", ["işe giriş", "ise giris", "giriş tarihi", "başlangıç", "start date", "işe başlama", "işe giriş tarihi"]),
        "iban": ("IBAN", False, "str", ["iban"])}},
}
# Sistem bazlı özel başlık eşlemeleri (alias listesine ek olarak öncelikli)
SOURCE_HINTS: Dict[str, Dict[str, Dict[str, str]]] = {
    "bizimhesap": {"contacts": {"cari ünvanı": "name", "vergi no / tc": "tax_number_or_id", "vergi no/tc": "tax_number_or_id", "cari kodu": "category", "borç": "balance"},
                   "products": {"stok kodu": "sku", "stok adı": "name", "satış fiyatı (kdv dahil)": "sale_price", "alış fiyatı (kdv dahil)": "purchase_price", "kdv oranı": "vat_rate"},
                   "invoices": {"fatura no": "invoice_number", "cari ünvanı": "contact_name", "genel toplam": "grand_total"}},
    "parasut": {"contacts": {"unvan": "name", "vergi numarası": "tax_number_or_id", "kısa ad": "company_title", "bakiye": "balance"},
                "products": {"ürün adı": "name", "ürün kodu": "sku", "stok miktarı": "stock_quantity", "satış fiyatı": "sale_price", "alış fiyatı": "purchase_price"},
                "invoices": {"fatura numarası": "invoice_number", "müşteri": "contact_name", "toplam": "grand_total", "düzenleme tarihi": "issue_date", "vade tarihi": "due_date"}},
    "logo_isbasi": {"contacts": {"cari hesap ünvanı": "name", "cari hesap kodu": "category", "vergi kimlik no": "tax_number_or_id"},
                    "products": {"malzeme kodu": "sku", "malzeme adı": "name", "malzeme açıklaması": "name"}},
    "mikro": {"contacts": {"cari ünvanı": "name", "cari kodu": "category", "vergi numarası": "tax_number_or_id", "bakiye": "balance"},
              "products": {"stok kodu": "sku", "stok adı": "name", "barkod": "barcode", "satış fiyatı": "sale_price", "alış fiyatı": "purchase_price"}},
    "sovos": {"contacts": {"ünvan": "name", "vergi kimlik no": "tax_number_or_id", "vergi dairesi": "tax_office"},
              "invoices": {"fatura no": "invoice_number", "ettn": "invoice_number", "düzenleme tarihi": "issue_date", "genel toplam": "grand_total"}},
}


def init(db):
    global _db
    _db = db


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _norm(s: Any) -> str:
    return re.sub(r"\s+", " ", str(s or "").replace("İ", "i").replace("I", "ı")).strip().lower()


def _to_money(v: Any) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v))
    if not s:
        return 0.0
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".") if len(s.split(",")[-1]) <= 2 else s.replace(",", "")
    elif s.count(".") > 1 or (s.count(".") == 1 and len(s.split(".")[-1]) == 3):
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        raise ValueError(f"Sayı okunamadı: {v}")


def _to_date(v: Any) -> Optional[str]:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    raise ValueError(f"Tarih okunamadı: {v}")


def _cast(field_type: str, v: Any) -> Any:
    if field_type == "str":
        return str(v).strip() if v not in (None, "") else None
    if field_type == "money" or field_type == "float":
        return _to_money(v)
    if field_type == "int":
        return int(round(_to_money(v)))
    if field_type == "date":
        return _to_date(v)
    n = _norm(v)
    if field_type == "contact_type":
        return "supplier" if any(k in n for k in ("tedarik", "satıcı", "supplier", "alış")) else "both" if ("her iki" in n or "both" in n) else "customer"
    if field_type == "product_type":
        return "service" if ("hizmet" in n or "service" in n) else "raw_material" if ("hammadde" in n or "raw" in n) else "product"
    if field_type == "invoice_type":
        return "purchase" if any(k in n for k in ("alış", "alis", "purchase", "gelen", "alım")) else "sales"
    if field_type == "account_type":
        return "cash_box" if ("kasa" in n or "cash" in n or "nakit" in n) else "pos" if "pos" in n else "credit_card" if ("kredi" in n or "kart" in n) else "bank"
    if field_type == "payment_status":
        return "paid" if any(k in n for k in ("ödendi", "odendi", "paid", "tahsil", "kapalı", "kapali")) else "partially_paid" if "kısm" in n else "unpaid"
    return v


def suggest_mapping(entity: str, columns: List[str], source: str) -> Dict[str, Optional[str]]:
    fields = ENTITIES[entity]["fields"]
    hints = (SOURCE_HINTS.get(source) or {}).get(entity) or {}
    mapping: Dict[str, Optional[str]] = {}
    used = set()
    ncols = [(c, _norm(c)) for c in columns]
    for col, n in ncols:  # önce sistem ipuçları
        if n in hints and hints[n] not in mapping:
            mapping[hints[n]] = col; used.add(col)
    for f, (_label, _req, _t, aliases) in fields.items():
        if f in mapping:
            continue
        exact = next((c for c, n in ncols if c not in used and n in aliases), None)
        partial = exact or next((c for c, n in ncols if c not in used and any(a == n or (len(a) > 3 and a in n) for a in aliases)), None)
        mapping[f] = partial
        if partial:
            used.add(partial)
    return mapping


def products_from_table_bytes(filename: str, data: bytes) -> List[dict]:
    """Excel/CSV stok listesini sütun eşlemesiyle ürün kayıtlarına çevirir (AI gerekmez)."""
    header, body = _read_table(filename, data)
    mapping = suggest_mapping("products", header, "other")
    if not mapping.get("name"):
        return []
    fields = ENTITIES["products"]["fields"]
    out: List[dict] = []
    for r in body:
        row = {header[i]: (r[i] if i < len(r) else None) for i in range(len(header))}
        rec: Dict[str, Any] = {}
        for f, (_label, _req, t, _a) in fields.items():
            col = mapping.get(f)
            raw = row.get(col) if col else None
            if raw in (None, ""):
                continue
            try:
                rec[f] = _cast(t, raw)
            except ValueError:
                continue
        if rec.get("name"):
            out.append(rec)
    return out


def _read_table(filename: str, data: bytes) -> tuple[List[str], List[List[Any]]]:
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|") if text.strip() else csv.excel
        rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(str(c).strip() for c in r)]
    else:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(status_code=400, detail="Dosya okunamadı. .xlsx veya .csv yükleyin (eski .xls dosyalarını Excel'de .xlsx olarak kaydedin).")
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True) if r and any(c not in (None, "") for c in r)]
    if not rows:
        raise HTTPException(status_code=400, detail="Dosya boş.")
    header_idx = next((i for i, r in enumerate(rows[:10]) if sum(1 for c in r if isinstance(c, str) and c.strip()) >= max(2, len([c for c in r if c not in (None, "")]) * 0.6)), 0)
    header = [str(c).strip() if c not in (None, "") else f"Sütun {i + 1}" for i, c in enumerate(rows[header_idx])]
    body = [list(r) + [None] * (len(header) - len(r)) for r in rows[header_idx + 1:]]
    return header, body


@router.get("/migration/sources")
async def migration_sources():
    return {"sources": SOURCES, "entities": [{"key": k, "label": v["label"], "fields": [{"key": f, "label": d[0], "required": d[1], "type": d[2]} for f, d in v["fields"].items()]} for k, v in ENTITIES.items()]}


@router.get("/migration/template/{entity}")
async def migration_template(entity: str):
    if entity not in ENTITIES:
        raise HTTPException(status_code=404, detail="Veri türü bulunamadı.")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = re.sub(r"[\\/*?:\[\]]", "-", ENTITIES[entity]["label"])[:30]
    ws.append([d[0] for d in ENTITIES[entity]["fields"].values()])
    sample = {"contacts": ["Örnek Müşteri A.Ş.", "customer", "1234567890", "Kadıköy", "info@ornek.com", "05321234567", "Örnek Mah. No:1", "İstanbul", "Kadıköy", "1500", "Genel", "", ""],
              "products": ["Örnek Ürün", "SKU-001", "8690000000001", "Genel", "Adet", "20", "100", "150", "25", "5", "product"],
              "invoices": ["FTR2026000001", "sales", "Örnek Müşteri A.Ş.", "1234567890", "2026-01-15", "2026-02-15", "1000", "200", "1200", "0", "unpaid", ""],
              "bank_accounts": ["Ana Kasa", "cash_box", "", "", "", "TRY", "5000"],
              "employees": ["Ali Veli", "12345678901", "Satış", "Temsilci", "05321234567", "ali@ornek.com", "30000", "40000", "2024-01-01", ""]}
    ws.append(sample.get(entity, []))
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="sablon_{entity}.xlsx"'})


@router.post("/migration/parse")
async def migration_parse(file: UploadFile = File(...), entity: str = Form(...), source: str = Form("other"), company_id: str = Form("comp_nexus_main_01")):
    if entity not in ENTITIES:
        raise HTTPException(status_code=400, detail="Geçersiz veri türü.")
    data = await file.read()
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Dosya en fazla 15 MB olabilir.")
    header, body = _read_table(file.filename, data)
    rows = [{header[i]: (c.isoformat() if isinstance(c, (datetime, date)) else c) for i, c in enumerate(r[:len(header)])} for r in body]
    upload = {"_id": str(uuid.uuid4()), "company_id": company_id, "entity": entity, "source": source, "filename": file.filename, "columns": header, "rows": rows, "created_at": _now()}
    await _db.migration_uploads.insert_one(upload)
    return {"upload_id": upload["_id"], "columns": header, "row_count": len(rows), "sample": rows[:8], "suggested_mapping": suggest_mapping(entity, header, source),
            "fields": [{"key": f, "label": d[0], "required": d[1], "type": d[2]} for f, d in ENTITIES[entity]["fields"].items()]}


async def _transform(upload: dict, mapping: Dict[str, Optional[str]]):
    entity = upload["entity"]; fields = ENTITIES[entity]["fields"]; company_id = upload["company_id"]
    key_field = ENTITIES[entity]["key"]
    coll = _db[ENTITIES[entity]["collection"]]
    missing_req = [fields[f][0] for f, d in fields.items() if d[1] and not mapping.get(f)]
    if missing_req:
        raise HTTPException(status_code=400, detail=f"Zorunlu alanlar eşlenmedi: {', '.join(missing_req)}")
    records, errors, seen = [], [], set()
    existing_keys: Dict[str, str] = {}
    if key_field:
        async for d in coll.find({"company_id": company_id}, {key_field: 1, "name": 1, "account_name": 1, "full_name": 1}):
            kv = _norm(d.get(key_field))
            if kv:
                existing_keys[kv] = d["_id"]
            nm = _norm(d.get("name") or d.get("account_name") or d.get("full_name"))
            if nm and key_field in ("tax_number_or_id", "tc_kimlik", "sku"):
                existing_keys.setdefault(f"name::{nm}", d["_id"])
    for i, row in enumerate(upload["rows"], start=2):
        rec: Dict[str, Any] = {}
        row_err = []
        for f, (label, req, t, _a) in fields.items():
            col = mapping.get(f)
            raw = row.get(col) if col else None
            try:
                val = _cast(t, raw) if raw not in (None, "") else None
            except ValueError as e:
                row_err.append(f"{label}: {e}"); val = None
            if req and val in (None, "") and t not in ("money", "float", "int"):
                row_err.append(f"{label} boş")
            if val is not None:
                rec[f] = val
        if entity == "invoices" and "grand_total" in rec and rec["grand_total"] <= 0:
            row_err.append("Genel toplam sıfır/negatif")
        kv = _norm(rec.get(key_field)) if key_field else ""
        nm = _norm(rec.get("name") or rec.get("account_name") or rec.get("full_name"))
        dup_key = kv or (f"name::{nm}" if nm else f"row::{i}")
        if dup_key in seen:
            row_err.append("Dosya içinde mükerrer kayıt")
        seen.add(dup_key)
        existing_id = existing_keys.get(kv) if kv else existing_keys.get(f"name::{nm}") if nm else None
        if entity == "invoices" and kv:
            existing_id = existing_keys.get(kv)
        records.append({"row": i, "data": rec, "existing_id": existing_id, "errors": row_err, "label": rec.get("name") or rec.get("account_name") or rec.get("full_name") or rec.get("invoice_number") or f"Satır {i}"})
        if row_err:
            errors.append({"row": i, "label": records[-1]["label"], "errors": row_err})
    return records, errors


@router.post("/migration/ai-map")
async def migration_ai_map(req: Dict[str, Any]):
    """Claude ile sütun eşlemesi öner (otomatik eşleme yetersiz kaldığında)."""
    upload = await _db.migration_uploads.find_one({"_id": req.get("upload_id")})
    if not upload:
        raise HTTPException(status_code=404, detail="Yükleme bulunamadı; dosyayı tekrar yükleyin.")
    from ai_service import ai_map_columns
    ent = ENTITIES[upload["entity"]]
    fields = [{"key": f, "label": d[0], "required": d[1]} for f, d in ent["fields"].items()]
    try:
        res = await ai_map_columns(ent["label"], fields, upload["columns"], upload["rows"][:5])
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"AI eşleme başarısız: {str(e)[:140]}")
    auto = suggest_mapping(upload["entity"], upload["columns"], upload["source"])
    merged = {f: res["mapping"].get(f) or auto.get(f) for f in ent["fields"]}
    return {"mapping": merged, "ai_mapping": res["mapping"], "notes": res["notes"], "mapped": sum(1 for v in merged.values() if v)}


@router.post("/migration/preview")
async def migration_preview(req: Dict[str, Any]):
    upload = await _db.migration_uploads.find_one({"_id": req.get("upload_id")})
    if not upload:
        raise HTTPException(status_code=404, detail="Yükleme bulunamadı; dosyayı tekrar yükleyin.")
    records, errors = await _transform(upload, req.get("mapping") or {})
    valid = [r for r in records if not r["errors"]]
    return {"total": len(records), "valid": len(valid), "invalid": len(errors), "new": sum(1 for r in valid if not r["existing_id"]), "existing": sum(1 for r in valid if r["existing_id"]),
            "errors": errors[:200], "sample": [{"row": r["row"], "label": r["label"], "status": "hata" if r["errors"] else ("güncelle" if r["existing_id"] else "yeni"), "data": r["data"]} for r in records[:30]]}


def _build_doc(entity: str, company_id: str, data: dict, batch_id: str) -> dict:
    base = {"_id": str(uuid.uuid4()), "company_id": company_id, "import_batch_id": batch_id, "source": "migration", "created_at": _now()}
    if entity == "contacts":
        return {**base, "type": "customer", "tax_number_or_id": "", "balance": 0.0, "credit_limit": 0.0, "category": "Genel", "is_e_invoice_user": False, "payment_term_days": 0, "late_fee_rate": 0.0, "b2b_enabled": False, "b2b_discount": 0.0, **data}
    if entity == "products":
        sku = data.get("sku") or f"MIG-{uuid.uuid4().hex[:6].upper()}"
        return {**base, "type": "product", "category": "Genel", "unit": "Adet", "vat_rate": 20, "purchase_price": 0.0, "sale_price": 0.0, "currency": "TRY", "stock_quantity": 0.0, "min_stock_alert": 5.0, "warehouse_id": "main_warehouse",
                "has_variants": False, "variant_options": [], "variants": [], "images": [], "show_in_b2b": True, "track_stock": True, "purchase_vat_rate": 20.0, "price_includes_vat": False, "tags": [], "is_active": True, **data, "sku": sku,
                "barcode": data.get("barcode") or f"868{str(uuid.uuid4().int)[:10]}"}
    if entity == "bank_accounts":
        return {**base, "type": "bank", "bank_name": data.get("bank_name") or ("Kasa" if data.get("type") == "cash_box" else ""), "currency": "TRY", "current_balance": 0.0, "pos_commission_rate": 0.0, **data}
    if entity == "employees":
        return {**base, "tc_kimlik": "", "department": "Genel", "position": "", "phone": "", "email": "", "salary": 0.0, "start_date": _now()[:10], "status": "active", "annual_leave_days": 14, "used_leave_days": 0, **data}
    return {**base, **data}


async def _invoice_doc(company_id: str, data: dict, batch_id: str) -> dict:
    contact = None
    if data.get("contact_tax_id"):
        contact = await _db.contacts.find_one({"company_id": company_id, "tax_number_or_id": data["contact_tax_id"]})
    if not contact and data.get("contact_name"):
        contact = await _db.contacts.find_one({"company_id": company_id, "name": {"$regex": f"^{re.escape(data['contact_name'])}$", "$options": "i"}})
    if not contact:
        contact = {"_id": str(uuid.uuid4()), "company_id": company_id, "type": "supplier" if data.get("invoice_type") == "purchase" else "customer", "name": data["contact_name"], "tax_number_or_id": data.get("contact_tax_id") or "", "balance": 0.0,
                   "category": "Aktarım", "is_e_invoice_user": False, "import_batch_id": batch_id, "source": "migration", "created_at": _now()}
        await _db.contacts.insert_one(contact)
    gt = float(data.get("grand_total") or 0)
    sub = float(data.get("subtotal") or 0) or round(gt / 1.2, 2)
    vat = float(data.get("vat_total") or 0) or round(gt - sub, 2)
    paid = float(data.get("paid_amount") or 0)
    ps = data.get("payment_status") or ("paid" if paid >= gt - 0.01 else "partially_paid" if paid > 0 else "unpaid")
    if ps == "paid" and paid <= 0:
        paid = gt
    return {"_id": str(uuid.uuid4()), "company_id": company_id, "invoice_number": data["invoice_number"], "invoice_type": data.get("invoice_type") or "sales", "e_type": "paper", "contact_id": contact["_id"], "contact_name": contact.get("name"),
            "contact_tax_id": contact.get("tax_number_or_id"), "issue_date": data["issue_date"], "due_date": data.get("due_date") or data["issue_date"], "items": [{"name": data.get("notes") or "Aktarılan fatura (devir)", "quantity": 1, "unit": "Adet", "unit_price": sub, "vat_rate": round(vat / sub * 100) if sub else 20, "discount_rate": 0, "total": sub, "vat_amount": vat}],
            "subtotal": sub, "vat_total": vat, "discount_total": 0.0, "grand_total": gt, "currency": "TRY", "status": "approved", "gib_status": None, "payment_status": ps, "paid_amount": paid, "notes": data.get("notes") or "", "source": "migration",
            "import_batch_id": batch_id, "created_at": _now()}


@router.post("/migration/import")
async def migration_import(req: Dict[str, Any]):
    upload = await _db.migration_uploads.find_one({"_id": req.get("upload_id")})
    if not upload:
        raise HTTPException(status_code=404, detail="Yükleme bulunamadı; dosyayı tekrar yükleyin.")
    entity, company_id = upload["entity"], upload["company_id"]
    on_dup = req.get("on_duplicate") if req.get("on_duplicate") in ("update", "skip") else "skip"
    records, errors = await _transform(upload, req.get("mapping") or {})
    coll = _db[ENTITIES[entity]["collection"]]
    batch = {"_id": str(uuid.uuid4()), "company_id": company_id, "entity": entity, "entity_label": ENTITIES[entity]["label"], "source": upload["source"], "filename": upload["filename"], "mapping": req.get("mapping"), "on_duplicate": on_dup,
             "inserted_ids": [], "updated": [], "skipped": 0, "failed": len(errors), "errors": errors[:500], "status": "done", "created_at": _now()}
    for r in records:
        if r["errors"]:
            continue
        if r["existing_id"]:
            if on_dup == "skip":
                batch["skipped"] += 1; continue
            prev = await coll.find_one({"_id": r["existing_id"]})
            upd = {k: v for k, v in r["data"].items() if k not in ("balance", "current_balance", "stock_quantity")} if entity in ("contacts", "bank_accounts", "products") else r["data"]
            await coll.update_one({"_id": r["existing_id"]}, {"$set": upd})
            batch["updated"].append({"id": r["existing_id"], "prev": {k: prev.get(k) for k in upd}})
            continue
        doc = await _invoice_doc(company_id, r["data"], batch["_id"]) if entity == "invoices" else _build_doc(entity, company_id, r["data"], batch["_id"])
        await coll.insert_one(doc)
        batch["inserted_ids"].append(doc["_id"])
    await _db.migration_batches.insert_one(batch)
    await _db.migration_uploads.delete_one({"_id": upload["_id"]})
    n_ins, n_upd = len(batch["inserted_ids"]), len(batch["updated"])
    return {"status": "success", "batch_id": batch["_id"], "inserted": n_ins, "updated": n_upd, "skipped": batch["skipped"], "failed": batch["failed"],
            "message": f"{ENTITIES[entity]['label']}: {n_ins} yeni, {n_upd} güncellendi, {batch['skipped']} atlandı, {batch['failed']} hatalı."}


def _batch_view(b: dict) -> dict:
    return {"id": b["_id"], "entity": b["entity"], "entity_label": b.get("entity_label"), "source": b.get("source"), "filename": b.get("filename"), "inserted": len(b.get("inserted_ids") or []), "updated": len(b.get("updated") or []),
            "skipped": b.get("skipped", 0), "failed": b.get("failed", 0), "status": b.get("status"), "created_at": b.get("created_at"), "rolled_back_at": b.get("rolled_back_at"), "error_count": len(b.get("errors") or [])}


@router.get("/migration/batches")
async def migration_batches(company_id: str = "comp_nexus_main_01"):
    return [_batch_view(b) for b in await _db.migration_batches.find({"company_id": company_id}).sort("created_at", -1).to_list(200)]


@router.get("/migration/batches/{batch_id}/errors.csv")
async def migration_batch_errors(batch_id: str):
    b = await _db.migration_batches.find_one({"_id": batch_id})
    if not b:
        raise HTTPException(status_code=404, detail="Parti bulunamadı.")
    buf = io.StringIO(); w = csv.writer(buf, delimiter=";")
    w.writerow(["Satır", "Kayıt", "Hatalar"])
    for e in b.get("errors") or []:
        w.writerow([e["row"], e["label"], " | ".join(e["errors"])])
    return StreamingResponse(io.BytesIO(("\ufeff" + buf.getvalue()).encode("utf-8")), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="aktarim_hatalar_{batch_id[:8]}.csv"'})


@router.post("/migration/batches/{batch_id}/rollback")
async def migration_rollback(batch_id: str):
    b = await _db.migration_batches.find_one({"_id": batch_id})
    if not b:
        raise HTTPException(status_code=404, detail="Parti bulunamadı.")
    if b.get("status") == "rolled_back":
        raise HTTPException(status_code=400, detail="Bu parti zaten geri alınmış.")
    coll = _db[ENTITIES[b["entity"]]["collection"]]
    r = await coll.delete_many({"_id": {"$in": b.get("inserted_ids") or []}})
    extra = 0
    if b["entity"] == "invoices":
        extra = (await _db.contacts.delete_many({"import_batch_id": batch_id})).deleted_count
    for u in b.get("updated") or []:
        await coll.update_one({"_id": u["id"]}, {"$set": u.get("prev") or {}})
    await _db.migration_batches.update_one({"_id": batch_id}, {"$set": {"status": "rolled_back", "rolled_back_at": _now()}})
    return {"status": "success", "deleted": r.deleted_count, "reverted": len(b.get("updated") or []), "message": f"{r.deleted_count} kayıt silindi, {len(b.get('updated') or [])} güncelleme geri alındı" + (f", {extra} otomatik cari silindi" if extra else "") + "."}


# ---------------- BizimHesap API ----------------
def _bh_headers(token: str, firm_id: str = "") -> Dict[str, str]:
    h = {
        "Key": BIZIMHESAP_KEY,
        "Token": token,
        "Accept": "application/json",
        "User-Agent": "TamKobi/1.0 (BizimHesap entegrasyonu)",
    }
    if firm_id:
        h["FirmId"] = firm_id
    return h


def _unwrap(payload: Any) -> List[dict]:
    """BizimHesap zarfı: {resultCode, errorText, data:{products|warehouses|inventory:[...]}}"""
    if isinstance(payload, list):
        return [x for x in payload if isinstance(x, dict)]
    if isinstance(payload, dict):
        if payload.get("resultCode") not in (None, 1, "1", True) and payload.get("errorText"):
            raise HTTPException(status_code=502, detail=f"BizimHesap hata: {payload.get('errorText')}")
        for k in ("data", "Data", "result"):
            if isinstance(payload.get(k), (dict, list)):
                return _unwrap(payload[k])
        for k in ("products", "items", "warehouses", "inventory", "customers", "Products", "Items"):
            if isinstance(payload.get(k), list):
                return [x for x in payload[k] if isinstance(x, dict)]
        return []
    return []


def _pick(d: dict, *keys, default=None):
    low = {str(k).lower(): v for k, v in d.items()}
    for k in keys:
        if k.lower() in low and low[k.lower()] not in (None, ""):
            return low[k.lower()]
    return default


def _bh_photo_url(product: dict) -> Optional[str]:
    """BizimHesap ürün görseli — farklı alan adları ve göreli yollar."""
    raw = _pick(
        product or {},
        "photo", "photoUrl", "photoURL", "image", "imageUrl", "imageURL",
        "picture", "thumbnail", "thumbnailUrl", "img", "productImage",
        default="",
    )
    if isinstance(raw, dict):
        raw = _pick(raw, "url", "src", "href", "path", default="") or ""
    if isinstance(raw, list) and raw:
        first = raw[0]
        raw = first.get("url") if isinstance(first, dict) else first
    url = str(raw or "").strip()
    if not url or url.lower() in ("null", "none", "undefined"):
        return None
    if url.startswith("//"):
        return "https:" + url
    if url.startswith("http://") or url.startswith("https://") or url.startswith("data:"):
        return url
    # Göreli yol → BizimHesap kökü
    if url.startswith("/"):
        return "https://bizimhesap.com" + url
    if url.startswith("uploads/") or url.startswith("files/") or url.startswith("images/"):
        return "https://bizimhesap.com/" + url
    return None


def _bh_http_detail(path: str, r: httpx.Response) -> str:
    snippet = (r.text or "").strip().replace("\n", " ")[:180]
    try:
        body = r.json()
        if isinstance(body, dict):
            snippet = str(body.get("errorText") or body.get("Message") or body.get("message") or snippet)
    except ValueError:
        pass
    return snippet


async def _bh_get(path: str, token: str, firm_id: str = "", timeout: float = 40.0) -> Any:
    try:
        async with httpx.AsyncClient(base_url=BIZIMHESAP_BASE, timeout=httpx.Timeout(timeout, connect=12.0), follow_redirects=True) as client:
            r = await client.get(path, headers=_bh_headers(token, firm_id))
    except httpx.ConnectTimeout:
        raise HTTPException(status_code=502, detail="BizimHesap sunucusuna bağlanılamadı (zaman aşımı). Sunucunun internet çıkışını kontrol edin.")
    except httpx.ConnectError as e:
        raise HTTPException(status_code=502, detail=f"BizimHesap sunucusuna bağlanılamadı: {e}")
    except httpx.ReadTimeout:
        raise HTTPException(status_code=502, detail=f"BizimHesap {path} yanıt vermedi (zaman aşımı). Katalog büyükse birkaç dakika sonra tekrar deneyin.")
    except httpx.HTTPError as e:
        raise HTTPException(status_code=502, detail=f"BizimHesap bağlantı hatası: {type(e).__name__}")
    if r.status_code in (401, 403):
        extra = _bh_http_detail(path, r)
        raise HTTPException(status_code=400, detail=f"BizimHesap token veya Firma ID geçersiz ({r.status_code}). {extra}".strip())
    if r.status_code == 429:
        raise HTTPException(status_code=429, detail="BizimHesap istek sınırı aşıldı (429). 1-2 dakika bekleyip tekrar deneyin.")
    if r.status_code >= 400:
        extra = _bh_http_detail(path, r)
        raise HTTPException(status_code=502, detail=f"BizimHesap {path} hatası: HTTP {r.status_code}" + (f" — {extra}" if extra else ""))
    try:
        return r.json()
    except ValueError:
        raise HTTPException(status_code=502, detail="BizimHesap yanıtı JSON değil.")


@router.get("/migration/bizimhesap/config")
async def bh_get_config(company_id: str = "comp_nexus_main_01"):
    c = await _db.migration_api_configs.find_one({"company_id": company_id, "provider": "bizimhesap"}) or {}
    return {"configured": bool(c.get("token_enc")), "firm_id": c.get("firm_id", ""), "last_test": c.get("last_test"), "last_import": c.get("last_import"), "last_customer_import": c.get("last_customer_import"), "token_mask": ("••••" + c["token_tail"]) if c.get("token_tail") else ""}


@router.put("/migration/bizimhesap/config")
async def bh_put_config(req: Dict[str, Any]):
    company_id = req.get("company_id", "comp_nexus_main_01")
    upd = {"firm_id": (req.get("firm_id") or "").strip(), "updated_at": _now()}
    if req.get("token"):
        tok = req["token"].strip()
        upd["token_enc"] = comm_service.encrypt(tok); upd["token_tail"] = tok[-4:]
    await _db.migration_api_configs.update_one({"company_id": company_id, "provider": "bizimhesap"}, {"$set": upd, "$setOnInsert": {"_id": str(uuid.uuid4()), "company_id": company_id, "provider": "bizimhesap"}}, upsert=True)
    return await bh_get_config(company_id)


async def _bh_creds(company_id: str, token: str = "", firm_id: str = "") -> tuple[str, str]:
    c = await _db.migration_api_configs.find_one({"company_id": company_id, "provider": "bizimhesap"}) or {}
    tok = (token or "").strip()
    fid = (firm_id or "").strip() or str(c.get("firm_id") or "").strip()
    if not tok:
        if not c.get("token_enc"):
            raise HTTPException(status_code=400, detail="Önce BizimHesap token'ını kaydedin.")
        try:
            tok = comm_service.decrypt(c["token_enc"])
        except Exception:
            raise HTTPException(status_code=400, detail="Kayıtlı token çözülemedi. Token'ı yeniden kaydedin.")
    return tok, fid


async def _bh_token(company_id: str) -> str:
    tok, _fid = await _bh_creds(company_id)
    return tok


async def _bh_products(company_id: str, token: str, firm_id: str = "", force: bool = False) -> tuple[List[dict], bool]:
    """/products ağır ve hız sınırlı → 15 dk önbellek. (ürünler, önbellekten_mi)"""
    key = {"company_id": company_id, "provider": "bizimhesap", "kind": "products"}
    cache = await _db.migration_api_cache.find_one(key)
    if cache and not force:
        try:
            age = (datetime.now(timezone.utc) - datetime.fromisoformat(cache["fetched_at"])).total_seconds()
        except (KeyError, ValueError):
            age = 1e9
        if age < 900:
            return cache["items"], True
    try:
        items = _unwrap(await _bh_get("/products", token, firm_id, timeout=90.0))
    except HTTPException as e:
        if e.status_code == 429 and cache:
            return cache["items"], True
        raise
    await _db.migration_api_cache.update_one(key, {"$set": {"items": items, "fetched_at": _now()}, "$setOnInsert": {"_id": str(uuid.uuid4()), **key}}, upsert=True)
    return items, False


@router.get("/migration/bizimhesap/warehouses")
async def bh_warehouses(company_id: str = "comp_nexus_main_01"):
    """Token kayıtlıysa depo listesini döner (ürün testi gerekmez)."""
    token, firm_id = await _bh_creds(company_id)
    warehouses = _unwrap(await _bh_get("/warehouses", token, firm_id, timeout=30.0))
    wh = [{"id": str(_pick(w, "id", "warehouseId", "depoId", "code", default="")), "name": _pick(w, "name", "title", "warehouseName", "depoAdi", default="Depo")} for w in warehouses]
    wh = [w for w in wh if w["id"]]
    return {"warehouses": wh, "count": len(wh)}


@router.post("/migration/bizimhesap/test")
async def bh_test(req: Dict[str, Any]):
    company_id = req.get("company_id", "comp_nexus_main_01")
    token, firm_id = await _bh_creds(company_id, req.get("token") or "", req.get("firm_id") or "")
    warehouses = _unwrap(await _bh_get("/warehouses", token, firm_id, timeout=30.0))
    wh = [{"id": str(_pick(w, "id", "warehouseId", "depoId", "code", default="")), "name": _pick(w, "name", "title", "warehouseName", "depoAdi", default="Depo")} for w in warehouses]
    product_error = None
    products: List[dict] = []
    cached = False
    try:
        products, cached = await _bh_products(company_id, token, firm_id, force=bool(req.get("force")))
    except HTTPException as e:
        product_error = e.detail
    active = sum(1 for p in products if str(p.get("isActive", 1)) in ("1", "True", "true"))
    with_photo = sum(1 for p in products if _bh_photo_url(p))
    res = {"ok": True, "cached": cached, "product_count": len(products), "active_count": active, "with_barcode": sum(1 for p in products if p.get("barcode")), "with_code": sum(1 for p in products if p.get("code")), "with_photo": with_photo, "warehouses": wh,
           "sample": [{k: p.get(k) for k in ("title", "code", "barcode", "price", "buyingPrice", "unit", "tax", "quantity", "category", "brand", "photo")} for p in products[:3]], "product_fields": sorted({k for p in products[:50] for k in p.keys()}),
           "product_error": product_error, "firm_id": firm_id}
    await _db.migration_api_configs.update_one({"company_id": company_id, "provider": "bizimhesap"}, {"$set": {"last_test": {"at": _now(), "product_count": len(products), "warehouse_count": len(wh), "with_photo": with_photo, "product_error": product_error}}})
    return res


@router.post("/migration/bizimhesap/import")
async def bh_import(req: Dict[str, Any]):
    company_id = req.get("company_id", "comp_nexus_main_01")
    token, firm_id = await _bh_creds(company_id)
    on_dup = req.get("on_duplicate") if req.get("on_duplicate") in ("update", "skip") else "update"
    with_images = req.get("with_images", True)
    products, _cached = await _bh_products(company_id, token, firm_id, force=bool(req.get("force")))
    stock: Dict[str, float] = {}
    warehouse_id = (req.get("warehouse_id") or "").strip() or None
    if req.get("with_stock", True) and warehouse_id:
        for it in _unwrap(await _bh_get(f"/inventory/{warehouse_id}", token, firm_id, timeout=90.0)):
            try:
                q = float(str(_pick(it, "qty", "quantity", "stock", "amount", default=0) or 0).replace(",", "."))
            except (TypeError, ValueError):
                continue
            for k in ("id", "productId", "product_id", "barcode", "code", "sku"):
                v = str(it.get(k) or "").strip().lower()
                if v:
                    stock[v] = q
    batch = {"_id": str(uuid.uuid4()), "company_id": company_id, "entity": "products", "entity_label": ENTITIES["products"]["label"], "source": "bizimhesap", "filename": "BizimHesap API /products", "on_duplicate": on_dup, "inserted_ids": [], "updated": [], "skipped": 0, "failed": 0, "errors": [], "status": "done", "created_at": _now(), "with_images": 0, "with_stock_qty": 0}
    only_active = req.get("only_active", True)
    for p in products:
        if only_active and str(p.get("isActive", 1)) not in ("1", "True", "true"):
            batch["skipped"] += 1; continue
        name = _pick(p, "title", "name", "productName")
        if not name:
            batch["failed"] += 1; batch["errors"].append({"row": 0, "label": str(_pick(p, "id", default="?")), "errors": ["Ürün adı (title) yok"]}); continue
        ext_id = str(_pick(p, "id", default="")).strip()
        code = str(_pick(p, "code", "sku", default="") or "").strip()
        barcode = str(_pick(p, "barcode", default="") or "").strip()
        sku = code or (f"BH-{ext_id[:8]}" if ext_id else None)
        data: Dict[str, Any] = {"name": str(name).strip()[:200], "sku": sku, "barcode": barcode or None, "category": str(p.get("category") or "").strip() or None, "unit": str(p.get("unit") or "").strip() or None,
                                "brand": str(p.get("brand") or "").strip() or None, "description": str(p.get("description") or "").strip() or None, "variant_name": str(p.get("variantName") or "").strip() or None, "is_active": True}
        for f, keys in (("sale_price", ("price", "salePrice")), ("purchase_price", ("buyingPrice", "purchasePrice")), ("vat_rate", ("tax", "taxRate", "vat"))):
            v = _pick(p, *keys)
            if v not in (None, ""):
                try:
                    data[f] = int(round(_to_money(v))) if f == "vat_rate" else _to_money(v)
                except ValueError:
                    pass
        if str(p.get("currency") or "").upper() in ("USD", "EUR", "GBP"):
            data["currency"] = str(p["currency"]).upper()
        if with_images:
            photo = _bh_photo_url(p)
            if photo:
                data["image_url"] = photo
                data["images"] = [photo]
                batch["with_images"] += 1
        qty = None
        for k in (ext_id.lower() if ext_id else None, barcode.lower() if barcode else None, code.lower() if code else None):
            if k and k in stock:
                qty = stock[k]; break
        if qty is None and p.get("quantity") not in (None, "") and not stock:
            try:
                qty = float(str(p["quantity"]).replace(",", "."))
            except (TypeError, ValueError):
                qty = None
        if qty is not None:
            data["stock_quantity"] = qty
            batch["with_stock_qty"] += 1
        data = {k: v for k, v in data.items() if v is not None}
        ors = ([{"bizimhesap_id": ext_id}] if ext_id else []) + ([{"barcode": barcode}] if barcode else []) + ([{"sku": code}] if code else [])
        existing = await _db.products.find_one({"company_id": company_id, "$or": ors}) if ors else None
        if existing:
            if on_dup == "skip":
                batch["skipped"] += 1; continue
            upd = {**data, "bizimhesap_id": ext_id}
            await _db.products.update_one({"_id": existing["_id"]}, {"$set": upd})
            batch["updated"].append({"id": existing["_id"], "prev": {k: existing.get(k) for k in upd}})
        else:
            doc = _build_doc("products", company_id, data, batch["_id"]); doc["bizimhesap_id"] = ext_id; doc["source"] = "bizimhesap"
            await _db.products.insert_one(doc); batch["inserted_ids"].append(doc["_id"])
    await _db.migration_batches.insert_one(batch)
    last = {"at": _now(), "inserted": len(batch["inserted_ids"]), "updated": len(batch["updated"]), "with_images": batch["with_images"], "with_stock_qty": batch["with_stock_qty"], "warehouse_id": warehouse_id}
    await _db.migration_api_configs.update_one({"company_id": company_id, "provider": "bizimhesap"}, {"$set": {"last_import": last}})
    parts = [f"BizimHesap: {len(products)} ürün okundu → {len(batch['inserted_ids'])} yeni stok kartı, {len(batch['updated'])} güncellendi, {batch['skipped']} atlandı"]
    if stock:
        parts.append(f"depodan {len(stock)} stok kaydı eşlendi ({batch['with_stock_qty']} karta yazıldı)")
    if with_images:
        parts.append(f"{batch['with_images']} üründe resim alındı")
    return {"status": "success", "batch_id": batch["_id"], "inserted": len(batch["inserted_ids"]), "updated": len(batch["updated"]), "skipped": batch["skipped"], "failed": batch["failed"], "with_stock": bool(stock), "with_images": batch["with_images"], "with_stock_qty": batch["with_stock_qty"],
            "message": "; ".join(parts) + "."}


_BH_TYPE_FIELDS = ("type", "customertype", "caritype", "kind", "accounttype", "carituru")
_BH_SUPPLIER_WORDS = ("tedarik", "satıcı", "satici", "supplier", "vendor", "alıcı_değil")


def bh_contact_type(row: dict, balance: float, forced: str = "auto") -> str:
    """BizimHesap cari kaydının müşteri mi tedarikçi mi olduğunu belirler.

    Eskiden herkes "customer" olarak kaydediliyordu (koşul `and False` ile kapatılmıştı),
    bu yüzden Cariler ekranındaki "Tedarikçiler" sekmesi boş kalıyordu. Sırasıyla:
    BizimHesap açıkça tür gönderiyorsa o, yoksa tek düzen hesap planındaki cari kodu
    (320… satıcılar, 120… alıcılar), o da yoksa bakiye işareti (borçlu olduğumuz cari
    tedarikçidir). Kullanıcı `contact_type` ile bu tahmini tamamen devre dışı bırakabilir.
    """
    if forced in ("customer", "supplier", "both"):
        return forced
    for key, value in (row or {}).items():
        if str(key).strip().lower().replace("_", "") not in _BH_TYPE_FIELDS:
            continue
        text = str(value or "").strip().lower()
        if any(w in text for w in _BH_SUPPLIER_WORDS):
            return "supplier"
        if "her iki" in text or "both" in text:
            return "both"
        if "müşteri" in text or "musteri" in text or "customer" in text or "alıcı" in text:
            return "customer"
    code = re.sub(r"\D", "", str((row or {}).get("code") or ""))
    if code.startswith("320") or code.startswith("329"):
        return "supplier"
    if code.startswith("120") or code.startswith("121"):
        return "customer"
    return "supplier" if balance < 0 else "customer"


@router.post("/migration/bizimhesap/import-customers")
async def bh_import_customers(req: Dict[str, Any]):
    """BizimHesap /customers → cariler (ünvan, VKN, vergi dairesi, telefon, e-posta, adres, yetkili, bakiye, çek/senet)."""
    company_id = req.get("company_id", "comp_nexus_main_01")
    token, firm_id = await _bh_creds(company_id)
    on_dup = req.get("on_duplicate") if req.get("on_duplicate") in ("update", "skip") else "update"
    invert = bool(req.get("invert_sign", False))
    only_bal = bool(req.get("only_with_balance", False))
    forced_type = req.get("contact_type") if req.get("contact_type") in ("customer", "supplier", "both") else "auto"
    customers = _unwrap(await _bh_get("/customers", token, firm_id, timeout=90.0))
    batch = {"_id": str(uuid.uuid4()), "company_id": company_id, "entity": "contacts", "entity_label": ENTITIES["contacts"]["label"], "source": "bizimhesap", "filename": "BizimHesap API /customers", "on_duplicate": on_dup, "inserted_ids": [], "updated": [], "skipped": 0, "failed": 0, "errors": [], "status": "done", "created_at": _now()}
    total_bal = 0.0
    types: Dict[str, int] = {}
    for c in customers:
        title = str(c.get("title") or "").strip()
        if not title:
            batch["failed"] += 1; batch["errors"].append({"row": 0, "label": str(c.get("id") or "?"), "errors": ["Ünvan boş"]}); continue
        try:
            bal = _to_money(c.get("balance")) * (-1 if invert else 1)
            cheque = _to_money(c.get("chequeandbond"))
        except ValueError:
            bal, cheque = 0.0, 0.0
        if only_bal and abs(bal) < 0.005:
            batch["skipped"] += 1; continue
        ext_id = str(c.get("id") or "").strip(); taxno = str(c.get("taxno") or "").strip()
        data = {k: v for k, v in {"name": title[:200], "tax_number_or_id": taxno or None, "tax_office": str(c.get("taxoffice") or "").strip() or None, "phone": str(c.get("phone") or "").strip() or None, "email": str(c.get("email") or "").strip().lower() or None,
                                  "address": str(c.get("address") or "").strip() or None, "contact_person": str(c.get("authorized") or "").strip() or None, "category": str(c.get("code") or "").strip() or None, "cheque_bond_balance": cheque or None,
                                  "currency": "TRY" if str(c.get("currency") or "TL").upper() in ("TL", "TRY") else str(c.get("currency")).upper()}.items() if v is not None}
        ctype = bh_contact_type(c, bal, forced_type)
        types[ctype] = types.get(ctype, 0) + 1
        ors = ([{"bizimhesap_id": ext_id}] if ext_id else []) + ([{"tax_number_or_id": taxno}] if taxno and taxno not in ("11111111111", "1111111111") else []) + [{"name": {"$regex": f"^{re.escape(title)}$", "$options": "i"}}]
        existing = await _db.contacts.find_one({"company_id": company_id, "$or": ors})
        total_bal += bal
        if existing:
            if on_dup == "skip":
                batch["skipped"] += 1; continue
            # Elle "Müşteri & Tedarikçi" yapılmış cariyi tek türe düşürmeyelim.
            upd = {**data, "bizimhesap_id": ext_id, "balance": bal, "opening_balance_source": "bizimhesap",
                   "type": "both" if existing.get("type") == "both" else ctype}
            await _db.contacts.update_one({"_id": existing["_id"]}, {"$set": upd})
            batch["updated"].append({"id": existing["_id"], "prev": {k: existing.get(k) for k in upd}})
        else:
            doc = _build_doc("contacts", company_id, {**data, "balance": bal, "type": ctype}, batch["_id"])
            doc.update({"bizimhesap_id": ext_id, "source": "bizimhesap", "tax_number_or_id": data.get("tax_number_or_id") or "", "opening_balance_source": "bizimhesap"})
            await _db.contacts.insert_one(doc); batch["inserted_ids"].append(doc["_id"])
    await _db.migration_batches.insert_one(batch)
    await _db.migration_api_configs.update_one({"company_id": company_id, "provider": "bizimhesap"}, {"$set": {"last_customer_import": {"at": _now(), "inserted": len(batch["inserted_ids"]), "updated": len(batch["updated"]), "total_balance": round(total_bal, 2), "types": types}}})
    breakdown = f"{types.get('customer', 0)} müşteri, {types.get('supplier', 0)} tedarikçi" + (f", {types['both']} müşteri & tedarikçi" if types.get("both") else "")
    return {"status": "success", "batch_id": batch["_id"], "read": len(customers), "inserted": len(batch["inserted_ids"]), "updated": len(batch["updated"]), "skipped": batch["skipped"], "failed": batch["failed"], "total_balance": round(total_bal, 2), "types": types,
            "message": f"BizimHesap: {len(customers)} cari okundu → {len(batch['inserted_ids'])} yeni, {len(batch['updated'])} güncellendi, {batch['skipped']} atlandı ({breakdown}). Toplam bakiye {total_bal:,.2f} ₺. (Aktarım Günlüğü'nden geri alınabilir.)"}
